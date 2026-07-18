"""
解析 26bct.docx 文档中的编程填空题并导入系统
用法: python import_fill_blanks.py
"""
import re
import json
import urllib.request
import urllib.error
from docx import Document

DOCX_PATH = r"c:\wza\xgpy-m\26bct.docx"
API_BASE = "http://localhost:3101/api"
TEACHER_USERNAME = "teacher"
TEACHER_PASSWORD = "meoo.local"

PLACEHOLDER_CHARS = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
placeholder_re = re.compile(r"[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]")


def split_answers(ans_text):
    """将答案文本分割为多个答案，处理方括号包裹的列表（列表内逗号不分割）"""
    ans_text = ans_text.strip()
    # 如果整个答案被方括号包裹，作为单个答案
    if ans_text.startswith('[') and ans_text.endswith(']'):
        return [ans_text] if ans_text else []
    # 否则用逗号分割，但要注意方括号内的逗号
    result = []
    current = []
    bracket_level = 0
    for ch in ans_text:
        if ch == '[':
            bracket_level += 1
            current.append(ch)
        elif ch == ']':
            bracket_level -= 1
            current.append(ch)
        elif ch == ',' and bracket_level == 0:
            part = ''.join(current).strip()
            if part:
                result.append(part)
            current = []
        else:
            current.append(ch)
    # 处理最后一个
    part = ''.join(current).strip()
    if part:
        result.append(part)
    return result


def extract_questions_from_doc(doc_path):
    """从文档中提取所有题目（题干+代码模板）"""
    doc = Document(doc_path)
    paragraphs = [p.text for p in doc.paragraphs]  # 保留原始空格

    questions = []
    current = None

    question_title_re = re.compile(r"^(\d+)、(.+)")
    # Python代码行起始模式
    code_start_re = re.compile(
        r"^(\s*)("
        r"def |for |if |while |print\s*\(|import |from |return |class |try |except |else:?|elif |#|"
        r"\w+\s*=|\w+\s*\(|\w+\s*\[|"
        r"[\"']|\.\w+|pass|break|continue|yield|with |raise |assert |global |nonlocal |lambda "
        r")"
    )

    for i, text in enumerate(paragraphs):
        stripped = text.strip()
        if stripped == "参考答案" or stripped.startswith("参考答案"):
            break

        m = question_title_re.match(stripped)
        if m:
            if current:
                questions.append(current)
            q_num = int(m.group(1))
            title = m.group(2).strip()
            current = {
                "num": q_num,
                "title": title,
                "all_lines": [],
            }
            continue

        if not current:
            continue
        current["all_lines"].append(text)

    if current:
        questions.append(current)

    # 后处理：每道题分离描述和代码
    for q in questions:
        lines = q["all_lines"]
        # 找到第一个代码行的索引
        code_start_idx = None
        for idx, line in enumerate(lines):
            stripped = line.strip()
            if not stripped:
                continue
            has_ph = bool(placeholder_re.search(stripped))
            # 去掉开头的占位符后再判断是否像代码
            stripped_no_ph = placeholder_re.sub("", stripped).strip()
            looks_like_code = bool(code_start_re.match(stripped_no_ph)) or stripped_no_ph.endswith(":")
            # 行以占位符开头，且去掉占位符后像代码 → 是代码
            if placeholder_re.match(stripped) and looks_like_code:
                code_start_idx = idx
                break
            if has_ph and looks_like_code:
                code_start_idx = idx
                break
            if looks_like_code and not re.match(r"^[请用打输运功该程算（一二三]", stripped):
                code_start_idx = idx
                break

        if code_start_idx is None:
            # 没找到代码行，全部当描述（异常情况）
            q["description"] = "\n".join([l.strip() for l in lines if l.strip()]).strip()
            q["blank_template"] = ""
        else:
            desc_lines = lines[:code_start_idx]
            code_lines = lines[code_start_idx:]
            # 描述部分：去除每行首尾空格，过滤空行
            q["description"] = "\n".join([l.strip() for l in desc_lines if l.strip()]).strip()
            # 代码部分：保留原始缩进，只移除首尾的纯空行
            while code_lines and not code_lines[0].strip():
                code_lines.pop(0)
            while code_lines and not code_lines[-1].strip():
                code_lines.pop()
            q["blank_template"] = "\n".join(code_lines)
        del q["all_lines"]

    return questions


def extract_answers_from_doc(doc_path):
    """从文档中提取所有题目的答案"""
    doc = Document(doc_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs]

    # 找到答案部分的开始
    answers_start = None
    for i, text in enumerate(paragraphs):
        if text.startswith("参考答案"):
            answers_start = i
            break

    if answers_start is None:
        return {}

    answers_section = paragraphs[answers_start:]

    answers_dict = {}
    current_num = None
    current_answers = []

    question_title_re = re.compile(r"^(\d+)、(.+)")
    answer_line_re = re.compile(r"^([①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳])[：:]\s*(.+)")

    for text in answers_section:
        if not text:
            continue
        if text == "一、编程填空题":
            continue

        m = question_title_re.match(text)
        if m:
            # 保存上一题答案
            if current_num is not None:
                answers_dict[current_num] = current_answers
            current_num = int(m.group(1))
            current_answers = []
            continue

        if text == "答案：":
            continue

        m = answer_line_re.match(text)
        if m and current_num is not None:
            ans_text = m.group(2).strip()
            ans_list = split_answers(ans_text)
            current_answers.append(ans_list)

    if current_num is not None and current_answers:
        answers_dict[current_num] = current_answers

    return answers_dict


def login(username, password):
    """登录获取session cookie"""
    data = json.dumps({"username": username, "password": password}).encode("utf-8")
    req = urllib.request.Request(
        f"{API_BASE}/auth/login",
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        resp = urllib.request.urlopen(req)
        cookie = resp.headers.get("Set-Cookie", "")
        body = json.loads(resp.read().decode("utf-8"))
        if body.get("data"):
            print(f"登录成功: {body['data'].get('real_name') or body['data'].get('username')}")
            return cookie
    except urllib.error.HTTPError as e:
        print(f"登录失败: {e.code} {e.read().decode('utf-8')}")
    return None


def create_task(cookie, task_data):
    """创建编程任务"""
    data = json.dumps(task_data, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        f"{API_BASE}/python/teacher/tasks",
        data=data,
        headers={
            "Content-Type": "application/json",
            "Cookie": cookie,
        },
        method="POST",
    )
    try:
        resp = urllib.request.urlopen(req)
        body = json.loads(resp.read().decode("utf-8"))
        return body.get("data")
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8")
        print(f"  创建失败: {e.code} {err_body}")
        return None


def main():
    print("=" * 60)
    print("解析 26bct.docx 文档...")
    print("=" * 60)

    questions = extract_questions_from_doc(DOCX_PATH)
    answers = extract_answers_from_doc(DOCX_PATH)

    print(f"\n提取到 {len(questions)} 道题目")
    print(f"提取到 {len(answers)} 题答案")

    # 打印前3题预览
    for q in questions[:3]:
        print(f"\n--- 第{q['num']}题: {q['title']} ---")
        print(f"描述: {q['description'][:80]}...")
        print(f"代码模板 ({len(q['blank_template'].splitlines())} 行):")
        print(q["blank_template"][:200])
        ans = answers.get(q["num"], [])
        print(f"答案: {ans}")

    # 验证题号连续性
    nums = sorted([q["num"] for q in questions])
    print(f"\n题号范围: {nums[0]} - {nums[-1]}")
    print(f"题号列表: {nums}")

    # 检查答案缺失
    missing = []
    for q in questions:
        if q["num"] not in answers:
            missing.append(q["num"])
    if missing:
        print(f"警告: 以下题目缺少答案: {missing}")
    else:
        print("所有题目都有答案 ✓")

    # 登录
    print("\n" + "=" * 60)
    print("登录系统...")
    print("=" * 60)
    cookie = login(TEACHER_USERNAME, TEACHER_PASSWORD)
    if not cookie:
        print("无法登录，退出")
        return

    # 从第3题开始导入（用户说前两题已创建）
    start_num = 3
    to_import = [q for q in questions if q["num"] >= start_num]
    print(f"\n将导入第 {start_num} - {nums[-1]} 题，共 {len(to_import)} 道题")

    confirm = input("\n确认导入? (y/n): ").strip().lower()
    if confirm != "y":
        print("已取消")
        return

    success = 0
    failed = []

    for q in to_import:
        num = q["num"]
        title = q["title"]
        print(f"\n正在导入第{num}题: {title}...")

        ans_list = answers.get(num, [])
        # 计算占位符数量
        ph_count = len(placeholder_re.findall(q["blank_template"]))
        # 如果答案数量不足，补充空列表
        while len(ans_list) < ph_count:
            ans_list.append([""])
        # 如果答案数量多了，截断
        ans_list = ans_list[:ph_count]

        # 每空默认分值 = 10 / 空数
        weight = round(10 / ph_count) if ph_count > 0 else 10
        weights = [weight] * ph_count
        # 调整最后一个使总分等于10
        if ph_count > 0:
            weights[-1] = 10 - sum(weights[:-1])

        task_data = {
            "title": f"第{num}题、{title}",
            "description": q["description"],
            "task_type": "fill_blank",
            "class_ids": [],
            "total_score": 10,
            "deadline": None,
            "allow_late_submission": False,
            "allow_run_code": True,
            "max_run_seconds": 3,
            "reference_code": "",
            "expected_output": "",
            "hint": "",
            "required_keywords": [],
            "blank_template": q["blank_template"],
            "blank_answers": ans_list,
            "blank_weights": weights,
        }

        result = create_task(cookie, task_data)
        if result:
            success += 1
            print(f"  ✓ 导入成功 (ID: {result['id']})")
        else:
            failed.append(num)
            print(f"  ✗ 导入失败")

    print("\n" + "=" * 60)
    print(f"导入完成: 成功 {success} 道，失败 {len(failed)} 道")
    if failed:
        print(f"失败题目: {failed}")
    print("=" * 60)


if __name__ == "__main__":
    main()
